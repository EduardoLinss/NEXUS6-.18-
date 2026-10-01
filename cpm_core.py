"""
cpm_core.py - Motor de analise do MIX / carga de trabalho por variante.

Substitui a logica de CPM4teste.py mantendo compatibilidade total:

    from cpm_core import (
        concatenar_planilhas_em_memoria,
        analisar_88_linhas_completas,   # continua devolvendo o dicionario antigo
        gerar_arquivo_concatenado,
        carregar_excel_sem_linhas_ocultas,
    )

Novidade principal:

    resultado = analisar_mix(df, popid, producao, absenteismo={"V1": 2, "V2": 1})

devolve um objeto `ResultadoAnalise` ja organizado por variante -> tarefas ->
percentual, que e exatamente o formato que a interface em colunas consome.
"""

from __future__ import annotations

import os
import logging
import re
import warnings
from copy import copy
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Iterable, Sequence

import pandas as pd
from openpyxl import Workbook, load_workbook

from setores import SETOR_PADRAO

logger = logging.getLogger(__name__)

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")


# ---------------------------------------------------------------------------
# 1. Constantes do calculo
# ---------------------------------------------------------------------------
# Formula original:  pct = (contagem * tempo) / ((producao * 300) / 100)
# O "/100" ja converte para porcentagem, por isso nao existe "* 100" no fim.
# Estes valores sao o PADRAO (FA2). Cada setor traz os seus — ver setores.py.
# Ficam aqui para o app desktop e para chamadas antigas continuarem valendo.
SEGUNDOS_DISPONIVEIS_POR_UNIDADE = 300
FATOR_PERCENTUAL = 100

# Absenteismo era (900 * faltas) / (producao * 3), que e a mesma formula acima
# com tempo = 900. Tratamos o absenteismo como "mais uma tarefa" da variante.
TEMPO_ABSENTEISMO = 900

# Nomes aceitos para a coluna de sequencia de producao (normalizados: sem
# espaco, ponto ou underscore, em maiuscula).
COLUNAS_SEQUENCIA = ("SEQ", "SEQUENCIA", "SEQUENCIAL", "ORDEM", "POSICAO", "NRSEQ")

# Horario em que o primeiro veiculo do intervalo passa pelo setor. E so o
# padrao do campo na tela — o valor usado vem sempre do formulario.
INICIO_TURNO_PADRAO = "07:05"
FIM_TURNO_PADRAO = "16:05"

# Pausas do turno (relogio parado para a producao). Entram apenas na estimativa
# de horario da linha do tempo — nao afetam a formula de carga.
PAUSAS_PADRAO = (
    ("09:00", "09:04"),
    ("10:55", "11:55"),
    ("14:00", "14:10"),
)

COMPRIMENTO_VEICULO_LONGO = 7400
COMPRIMENTO_CINTA_CARDAN = 8000


def _texto_celula(valor) -> str:
    """Celula em texto, sem o '.0' que o pandas cola em numero inteiro."""
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return ""
    if isinstance(valor, float) and float(valor).is_integer():
        return str(int(valor))
    return str(valor).strip()


def capacidade_disponivel(producao: float,
                          takt: float = SEGUNDOS_DISPONIVEIS_POR_UNIDADE) -> float:
    """Denominador comum de todos os percentuais de carga."""
    return (producao * takt) / FATOR_PERCENTUAL


def carga_percentual(contagem: float, tempo: float, producao: float,
                     takt: float = SEGUNDOS_DISPONIVEIS_POR_UNIDADE) -> float:
    capacidade = capacidade_disponivel(producao, takt)
    if capacidade <= 0:
        return 0.0
    return (contagem * tempo) / capacidade


# ---------------------------------------------------------------------------
# 2. Criterios de contagem (substituem as ~20 funcoes duplicadas)
# ---------------------------------------------------------------------------
class Criterio:
    """Devolve uma mascara booleana sobre o DataFrame filtrado."""

    def mascara(self, ctx: "ContextoDados") -> pd.Series:  # pragma: no cover
        raise NotImplementedError

    def colunas(self) -> tuple[str, ...]:  # pragma: no cover
        return ()


@dataclass(frozen=True)
class Contem(Criterio):
    """Texto contido na coluna (case-insensitive, igual ao comportamento antigo)."""

    coluna: str
    valor: str

    def mascara(self, ctx: "ContextoDados") -> pd.Series:
        serie = ctx.texto(self.coluna)
        if serie is None:
            return ctx.falso()
        return serie.str.contains(self.valor.upper(), regex=False, na=False)

    def colunas(self) -> tuple[str, ...]:
        return (self.coluna,)


@dataclass(frozen=True)
class MaiorQue(Criterio):
    """Coluna numerica acima de um limite (ex.: COMPRIM > 7400)."""

    coluna: str
    limite: float

    def mascara(self, ctx: "ContextoDados") -> pd.Series:
        serie = ctx.numero(self.coluna)
        if serie is None:
            return ctx.falso()
        return (serie > self.limite).fillna(False)

    def colunas(self) -> tuple[str, ...]:
        return (self.coluna,)


@dataclass(frozen=True)
class NaoContem(Criterio):
    """Texto ausente da coluna (ex.: OPCIONAIS sem 'DISCO')."""

    coluna: str
    valor: str

    def mascara(self, ctx: "ContextoDados") -> pd.Series:
        serie = ctx.texto(self.coluna)
        if serie is None:
            return ctx.falso()
        return ~serie.str.contains(self.valor.upper(), regex=False, na=False)

    def colunas(self) -> tuple[str, ...]:
        return (self.coluna,)


@dataclass(frozen=True)
class MenorQue(Criterio):
    """Coluna numerica abaixo de um limite."""

    coluna: str
    limite: float

    def mascara(self, ctx: "ContextoDados") -> pd.Series:
        serie = ctx.numero(self.coluna)
        if serie is None:
            return ctx.falso()
        return (serie < self.limite).fillna(False)

    def colunas(self) -> tuple[str, ...]:
        return (self.coluna,)


@dataclass(frozen=True)
class EmSequencia(Criterio):
    """
    Regra de sequencia: olha veiculos CONSECUTIVOS na ordem de producao.

    Exemplo: "3 ou mais veiculos seguidos com BARRA TRASEIRA = COM, a variante
    atua 1 vez" ->
        EmSequencia(Contem("BARRA TRASEIRA", "COM"), minimo=3, modo="primeiro")

    Devolve uma mascara, igual aos outros criterios, entao a contagem continua
    sendo `mascara.sum()` e a formula de carga nao muda. O que muda e QUAIS
    linhas ficam marcadas:

      primeiro  - marca a primeira linha de cada sequencia -> conta 1 por sequencia
      todos     - marca todos os veiculos das sequencias que qualificam
      excedente - marca so o que passa do minimo (o 3o, 4o... de cada sequencia)

    A ordem usada e a das linhas analisadas. Linhas ocultas e as filtradas por
    TVL nao entram, entao "consecutivo" aqui significa consecutivo NO RECORTE,
    nao necessariamente na linha fisica.
    """

    criterio: Criterio
    minimo: int = 3
    modo: str = "primeiro"

    def mascara(self, ctx: "ContextoDados") -> pd.Series:
        base = self.criterio.mascara(ctx).to_numpy(dtype=bool, copy=False)
        marcado = [False] * len(base)
        minimo = max(1, int(self.minimo))

        i = 0
        while i < len(base):
            if not base[i]:
                i += 1
                continue
            fim = i
            while fim < len(base) and base[fim]:
                fim += 1
            if (fim - i) >= minimo:
                if self.modo == "todos":
                    for k in range(i, fim):
                        marcado[k] = True
                elif self.modo == "excedente":
                    for k in range(i + minimo - 1, fim):
                        marcado[k] = True
                else:  # primeiro
                    marcado[i] = True
            i = fim

        return pd.Series(marcado, index=ctx.df.index)

    def colunas(self) -> tuple[str, ...]:
        return self.criterio.colunas()


@dataclass(frozen=True)
class EmJanela(Criterio):
    """
    Regra de janela: "N ou mais dentro de cada X veiculos", sem exigir que
    estejam colados.

    Exemplo: "a cada 6 veiculos, se 3 tiverem BARRA TRASEIRA = COM, a variante
    atua 1 vez" ->
        EmJanela(Contem("BARRA TRASEIRA", "COM"), janela=6, minimo=3)

    A contagem NAO se sobrepoe: quando uma janela dispara, os veiculos ate o
    gatilho sao consumidos e a varredura recomeca depois dele. Sem isso,
    6 veiculos seguidos contariam 4 vezes (uma por janela deslizante).

      primeiro - marca o veiculo que fecha o minimo -> 1 por disparo
      todos    - marca os `minimo` veiculos que fecharam a janela
    """

    criterio: Criterio
    janela: int = 6
    minimo: int = 3
    modo: str = "primeiro"

    def mascara(self, ctx: "ContextoDados") -> pd.Series:
        base = self.criterio.mascara(ctx).to_numpy(dtype=bool, copy=False)
        total = len(base)
        marcado = [False] * total
        janela = max(1, int(self.janela))
        minimo = max(1, int(self.minimo))

        i = 0
        while i < total:
            fim = min(i + janela, total)
            atingidos = [k for k in range(i, fim) if base[k]]
            if len(atingidos) >= minimo:
                gatilho = atingidos[minimo - 1]
                if self.modo == "todos":
                    for k in atingidos[:minimo]:
                        marcado[k] = True
                else:
                    marcado[gatilho] = True
                i = gatilho + 1
            else:
                i += 1

        return pd.Series(marcado, index=ctx.df.index)

    def colunas(self) -> tuple[str, ...]:
        return self.criterio.colunas()


@dataclass(frozen=True)
class E(Criterio):
    """Interseccao de criterios (AND)."""

    criterios: tuple[Criterio, ...]

    def __init__(self, *criterios: Criterio):
        object.__setattr__(self, "criterios", tuple(criterios))

    def mascara(self, ctx: "ContextoDados") -> pd.Series:
        resultado = None
        for criterio in self.criterios:
            atual = criterio.mascara(ctx)
            resultado = atual if resultado is None else (resultado & atual)
        return ctx.falso() if resultado is None else resultado

    def colunas(self) -> tuple[str, ...]:
        return tuple(c for crit in self.criterios for c in crit.colunas())


# ---------------------------------------------------------------------------
# 3. Catalogo declarativo: tudo o que antes era codigo agora e dado
# ---------------------------------------------------------------------------
@dataclass(frozen=True)
class Tarefa:
    nome: str            # rotulo exibido na coluna da interface
    tempo: float         # segundos por veiculo
    criterio: Criterio
    chave_legado: str = ""   # chave usada no dicionario antigo
    id: str = ""             # identificador estavel (usado na tela de edicao)


@dataclass(frozen=True)
class Variante:
    id: str
    nome: str
    tarefas: tuple[Tarefa, ...]


@dataclass(frozen=True)
class ContadorMix:
    chave: str           # chave do dicionario legado
    rotulo: str          # rotulo bonito para a interface
    criterio: Criterio


# --- MIX (apenas contagem, sem percentual) ---------------------------------
# Cada setor tem seus proprios itens contados no CPM MIX — ver
# `catalogo_mix_do_setor()` logo abaixo. CATALOGO_MIX_FA2 e o catalogo
# historico (FA2); CATALOGO_MIX_1_2 e o da FA1.2.
CATALOGO_MIX_FA2: tuple[ContadorMix, ...] = (
    ContadorMix("10x4", "10x4", Contem("MODELO", "10X4")),
    ContadorMix("8x4", "8x4", Contem("MODELO", "8X4")),
    ContadorMix("8x2", "8x2", Contem("MODELO", "8X2")),
    ContadorMix("6x6", "6x6", Contem("MODELO", "6X6")),
    ContadorMix("6x4", "6x4", Contem("MODELO", "6X4")),
    ContadorMix("6x2", "6x2", Contem("MODELO", "6X2")),
    ContadorMix("4x4", "4x4", Contem("MODELO", "4X4")),
    ContadorMix("4x2", "4x2", Contem("MODELO", "4X2")),
    ContadorMix("BARRA_TRASEIRA", "Barra traseira", Contem("BARRA TRASEIRA", "COM")),
    ContadorMix("ABAFADOR", "Abafador", Contem("OPCIONAIS", "ABAFADOR")),
    ContadorMix("HD_FRONT", "HD Front", Contem("OPCIONAIS", "HD FRONT")),
    ContadorMix("SUSP_Z", "Carros a mola", Contem("SUSP TRASEIRA", "Z")),
    ContadorMix("DISCO", "Disco", Contem("OPCIONAIS", "DISCO")),
    ContadorMix("SDAR", "SDAR", Contem("OPCIONAIS", "SDAR")),
    ContadorMix("LTDA", "LTDA", Contem("OPCIONAIS", "LTDA")),
    ContadorMix("FUP", "Viga FUP", Contem("OPCIONAIS", "FUP")),
    ContadorMix("RB", "RB885", Contem("RPX", "RB885")),
    ContadorMix("RBX", "RB756", Contem("RPX", "RB756")),
    ContadorMix("GUATEMALA", "GUATEMALA", Contem("MU", "GT01")),
    ContadorMix("6x2_DISCO", "6x2 disco",
                E(Contem("MODELO", "6X2"), Contem("OPCIONAIS", "DISCO"))),
    ContadorMix("4x2_Z", "4x2 Z",
                E(Contem("MODELO", "4X2"), Contem("SUSP TRASEIRA", "Z"))),
    ContadorMix("6X4HT", "6x4 HT",
                E(Contem("MODELO", "6X4"), Contem("OPCIONAIS", "BT400"))),
    ContadorMix("VEICULO_LONGO", "Veiculo longo",
                MaiorQue("COMPRIM", COMPRIMENTO_VEICULO_LONGO)),
)

# Compatibilidade: quem importava CATALOGO_MIX continua valendo (é o da FA2).
CATALOGO_MIX = CATALOGO_MIX_FA2

# Itens contados no CPM MIX da FA1.2. Os modelos (6x2/6x4/8x2/8x4/...) e
# "Veiculo longo" seguem a mesma logica da FA2. SDAR e LTDA, aqui, so entram
# via cruzamento com a MONA (ver TAGS_MONA_PADRAO_1_2 abaixo) — por isso nao
# aparecem como Contem("OPCIONAIS", ...) direto igual na FA2.
CATALOGO_MIX_1_2: tuple[ContadorMix, ...] = (
    ContadorMix("10x4", "10x4", Contem("MODELO", "10X4")),
    ContadorMix("8x4", "8x4", Contem("MODELO", "8X4")),
    ContadorMix("8x2", "8x2", Contem("MODELO", "8X2")),
    ContadorMix("6x6", "6x6", Contem("MODELO", "6X6")),
    ContadorMix("6x4", "6x4", Contem("MODELO", "6X4")),
    ContadorMix("6x2", "6x2", Contem("MODELO", "6X2")),
    ContadorMix("4x4", "4x4", Contem("MODELO", "4X4")),
    ContadorMix("4x2", "4x2", Contem("MODELO", "4X2")),
    ContadorMix("GAS", "Gás", Contem("GÁS", "GÁS")),
    ContadorMix("N_MODULE", "N_MODULE", Contem("OPCIONAIS", "N_MODULE")),
    ContadorMix("VEICULO_LONGO", "Veiculo longo",
                MaiorQue("COMPRIM", COMPRIMENTO_VEICULO_LONGO)),
)


def catalogo_mix_do_setor(setor: str | None = None) -> tuple[ContadorMix, ...]:
    """Itens contados no CPM MIX, por setor — mesmo padrão de
    `gargalos.colunas_regra_do_setor()`."""
    setor_id = setor or SETOR_PADRAO
    if setor_id == "1-2":
        return CATALOGO_MIX_1_2
    return CATALOGO_MIX_FA2


def _contador_texto_opcionais(texto: str) -> ContadorMix:
    """
    Monta um ContadorMix que so procura `texto` dentro de OPCIONAIS — usado
    tanto para as tags fixas da MONA (abaixo) quanto para o texto avulso que
    o usuario digita na tela ao cruzar com uma MONA fora do padrao POS_.
    """
    chave = "MONA_" + re.sub(r"\W+", "_", texto).strip("_").upper()
    return ContadorMix(chave, texto, Contem("OPCIONAIS", texto))


@dataclass(frozen=True)
class TagMona:
    """
    Uma variante especial que a MONA registra e que NAO existe no mix geral
    da fabrica — por isso precisa ser cruzada e contada a parte.

    texto   : o que entra em OPCIONAIS/Variantes quando o POPID tem essa tag,
              e o que o contador do CPM MIX procura (Contem("OPCIONAIS", ...)).
    padroes : expressao(oes) regex (sem diferenciar maiusculas/minusculas) que
              reconhecem essa tag pelo NOME DO ITEM numa planilha MONA no
              formato tecnico da Funcao 4 (POPID + colunas POS_, com o nome
              do item na 2a linha) — ex.: o item "SMOLA 658 / 660" bate com
              o padrao r"658\\D*660". So usado pra filtrar/mapear colunas
              dessa planilha; a lista simples de POPIDs nao passa por aqui.
    """
    texto: str
    padroes: tuple[str, ...] = ()


# Variantes que a MONA registra e que ja NAO existem como coluna no mix geral
# da fabrica — por isso sao cruzadas e contadas a parte no CPM MIX, tanto no
# Filtro MIX quanto na Analise. O resto das colunas de uma planilha MONA no
# formato da Funcao 4 (POS_) e ignorado, porque ja esta coberto pelo mix
# geral.
#
# Essa lista embutida e so o PONTO DE PARTIDA: a lista de verdade e editavel
# pela tela "/tags-mona" (por setor) e fica em tags_mona.json / tags_mona_1_2
# .json (ver tags_mona.py) — sem esse arquivo (primeira vez, ou fora de um
# deploy), usa-se esta aqui. TAGS_MONA_PADRAO e a da FA2 (nome mantido para
# nao quebrar quem ja importava direto).
TAGS_MONA_PADRAO: tuple[TagMona, ...] = (
    TagMona("MOLA 658|660", (r"658\D*660",)),
    TagMona("MOLA 649|650", (r"649\D*650",)),
    TagMona("CILINDRO ELEVADO", (r"CILINDRO\s*ELEV",)),
)

# Tags da MONA da FA1.2 — mesmo mecanismo, lista propria. "texto" e o rotulo
# que entra em OPCIONAIS/CPM MIX; "padroes" reconhece o item pelo nome real
# na planilha MONA (formato tecnico da Funcao 4).
TAGS_MONA_PADRAO_1_2: tuple[TagMona, ...] = (
    TagMona("LTDA", (r"LTDA",)),
    TagMona("CALHA INVERTIDA", (r"CALHA\s*INVERTIDA",)),
    TagMona("RADAR LATERAL", (r"RADAR\s*LATERAL",)),
    TagMona("SDAR", (r"SUSP\s*DIANT\s*ARX",)),
    TagMona("DIR. ELETRICA", (r"DIR\.?\s*ELETRICA",)),
    TagMona("BATERIA TRASEIRA", (r"BATERIA\s*TRASEI",)),
    TagMona("K1 20L", (r"K1\s*20\s*L",)),
    TagMona("CHICOTE FOU", (r"CHICOTE\s*FOU",)),
    TagMona('SOQUETE TANK "D', (r'SOQUETE\s*TANK\s*"?D',)),
)


def tags_mona_padrao_do_setor(setor: str | None = None) -> tuple[TagMona, ...]:
    """Lista embutida de tags da MONA, por setor (ponto de partida sem
    tags_mona.json — ver tags_mona_ativas())."""
    setor_id = setor or SETOR_PADRAO
    if setor_id == "1-2":
        return TAGS_MONA_PADRAO_1_2
    return TAGS_MONA_PADRAO


def tags_mona_ativas(setor: str | None = None) -> tuple[TagMona, ...]:
    """
    As tags da MONA em uso agora para o setor dado — lidas do
    tags_mona*.json daquele setor (editavel pela tela) quando existir; caso
    contrario, a lista embutida (`tags_mona_padrao_do_setor`). Chamada a cada
    analise/cruzamento, nunca cacheada, pra uma edicao pela tela valer na
    hora, sem reiniciar o servidor.
    """
    padrao = tags_mona_padrao_do_setor(setor)
    try:
        import tags_mona as tags_mona_mod
    except Exception as erro:  # noqa: BLE001
        logger.warning("tags_mona: modulo indisponivel, usando lista padrao (%s)", erro)
        return padrao
    try:
        return tags_mona_mod.carregar(setor)
    except Exception as erro:  # noqa: BLE001
        logger.warning("tags_mona: usando lista padrao (%s)", erro)
        return padrao


# --- Variantes -------------------------------------------------------------
CATALOGO_PADRAO: tuple[Variante, ...] = (
    Variante("V1", "Variante 1", (
        Tarefa("Haste traseira", 200, Contem("SUSP TRASEIRA", "Z"), "HASTE TRASEIRA v1"),
        Tarefa("HD Front", 200, Contem("OPCIONAIS", "HD FRONT"), "HD FRONT v1"),
        Tarefa("Abafador", 260, Contem("OPCIONAIS", "ABAFADOR"), "ABAFDOR v1"),
        Tarefa("SDAR", 60, Contem("OPCIONAIS", "SDAR"), "SDAR v1"),
        Tarefa("8x4", 223, Contem("MODELO", "8X4"), "8x4 v1"),
        Tarefa("8x2", 223, Contem("MODELO", "8X2"), "8x2 v1"),
    )),
    Variante("V2", "Variante 2", (
        Tarefa("Amarracao de disco", 290,
               E(Contem("MODELO", "6X2"), Contem("OPCIONAIS", "DISCO")),
               "amarracao de disco v2"),
        Tarefa("Barra traseira", 219, Contem("BARRA TRASEIRA", "COM"), "barra traseira v2"),
        Tarefa("LTDA", 120, Contem("OPCIONAIS", "LTDA"), "ltda v2"),
        Tarefa("8x4", 202, Contem("MODELO", "8X4"), "8x4 v2"),
        Tarefa("PM cardan", 120, Contem("MODELO", "8X2"), "pm cardan v2"),
        Tarefa("4x2 Z", 120,
               E(Contem("MODELO", "4X2"), Contem("SUSP TRASEIRA", "Z")), "4x2_Z v2"),
        Tarefa("Placa de ancoragem", 255,
               MaiorQue("COMPRIM", COMPRIMENTO_VEICULO_LONGO), "placa ancoragem v2"),
    )),
    Variante("V3", "Variante 3", (
        Tarefa("Flexivel de freio", 20, Contem("MODELO", "8X2"), "flexivel freio v3"),
        Tarefa("Barra traseira", 154, Contem("BARRA TRASEIRA", "COM"), "barra traseira v3"),
        Tarefa("Cinta", 178, MaiorQue("COMPRIM", COMPRIMENTO_VEICULO_LONGO), "cinta v3"),
        Tarefa("8x4", 267, Contem("MODELO", "8X4"), "8x4 v3"),
        Tarefa("Cinta cardan + abafador", 255,
               E(MaiorQue("COMPRIM", COMPRIMENTO_CINTA_CARDAN),
                 Contem("OPCIONAIS", "ABAFADOR")),
               "cinta cardan + abafador v3"),
    )),
    Variante("V4", "Variante 4", (
        Tarefa("8x4 2.1", 338, Contem("MODELO", "8X4"), "8x4 2.1 v4"),
        Tarefa("8x4 2.4", 267, Contem("MODELO", "8X4"), "8X4 2.4 v4"),
        Tarefa("Viga FUP", 190, Contem("OPCIONAIS", "FUP"), "fup v4"),
        Tarefa("Placa de ancoragem", 255,
               MaiorQue("COMPRIM", COMPRIMENTO_VEICULO_LONGO), "placa de ancoragem v4"),
    )),
)

# Compatibilidade: quem importava CATALOGO_VARIANTES continua funcionando.
CATALOGO_VARIANTES = CATALOGO_PADRAO


def catalogo_ativo(setor: str | None = None) -> tuple[Variante, ...]:
    """
    Catalogo de variantes/tarefas em uso PARA O SETOR DADO. Se existir
    `dados/<setor>/catalogo.json`, ele manda; caso contrario usa o embutido
    (`CATALOGO_PADRAO`). O import e tardio de proposito, para o modulo
    `catalogo` poder importar as classes daqui sem ciclo.

    Sem `setor`, cai no setor padrao (FA2) — por isso quem analisa outro
    setor (ex.: FA1.2) precisa sempre passar o `setor` aqui, senao as
    tarefas mostradas continuam sendo as da FA2.
    """
    try:
        from catalogo import carregar
        return carregar(setor)
    except Exception as erro:  # noqa: BLE001
        logger.warning("catalogo: usando catalogo padrao (%s)", erro)
        return CATALOGO_PADRAO


# Aliases das chaves antigas que tinham acento/grafia diferente, para que o
# codigo legado que le o dicionario continue funcionando sem alteracao.
ALIASES_LEGADO = {
    "amarracao de disco v2": "amarração de disco v2",
    "flexivel freio v3": "flexível freio v3",
}


# ---------------------------------------------------------------------------
# 4. Contexto de dados (normalizacao feita uma unica vez, com cache)
# ---------------------------------------------------------------------------
class ContextoDados:
    """
    Guarda o recorte analisado e devolve versoes normalizadas das colunas.

    Diferenca importante em relacao ao codigo antigo: nao se converte o
    DataFrame inteiro para string. Vazios continuam vazios (antes viravam
    a string "NAN", o que podia gerar match indevido) e colunas numericas
    continuam numericas.
    """

    def __init__(self, df: pd.DataFrame):
        self.df = df
        self._texto: dict[str, pd.Series] = {}
        self._numero: dict[str, pd.Series] = {}
        self._falso = pd.Series(False, index=df.index)
        self._coluna_seq: str | None | bool = False  # False = ainda nao procurou
        self.colunas_ausentes: set[str] = set()

    def texto(self, coluna: str) -> pd.Series | None:
        if coluna not in self.df.columns:
            self.colunas_ausentes.add(coluna)
            return None
        if coluna not in self._texto:
            self._texto[coluna] = (
                self.df[coluna].astype("object").map(
                    lambda v: "" if v is None or (isinstance(v, float) and pd.isna(v))
                    else str(v).strip().upper()
                )
            )
        return self._texto[coluna]

    def numero(self, coluna: str) -> pd.Series | None:
        if coluna not in self.df.columns:
            self.colunas_ausentes.add(coluna)
            return None
        if coluna not in self._numero:
            self._numero[coluna] = pd.to_numeric(self.df[coluna], errors="coerce")
        return self._numero[coluna]

    def falso(self) -> pd.Series:
        return self._falso

    def contar(self, criterio: Criterio) -> int:
        return int(criterio.mascara(self).sum())

    def identificar(self, criterio: Criterio) -> list[dict]:
        """
        POPID e SEQ dos veiculos que a regra marcou, para o detalhamento na
        tela. Em regra de sequencia/janela, sao os veiculos que dispararam.
        """
        mascara = criterio.mascara(self).to_numpy(dtype=bool, copy=False)
        coluna_seq = self.coluna_sequencia()

        itens = []
        for posicao, marcado in enumerate(mascara):
            if not marcado:
                continue
            linha = self.df.iloc[posicao]
            itens.append({
                "popid": _texto_celula(linha.get("POPID")),
                "seq": _texto_celula(linha.get(coluna_seq)) if coluna_seq else "",
                "posicao": posicao + 1,
            })
        return itens

    def identificar_todos(self) -> list[dict]:
        """POPID e SEQ de cada posição do recorte, na ordem de produção."""
        coluna_seq = self.coluna_sequencia()
        veiculos = []
        for posicao in range(len(self.df)):
            linha = self.df.iloc[posicao]
            veiculos.append({
                "popid": _texto_celula(linha.get("POPID")),
                "seq": _texto_celula(linha.get(coluna_seq)) if coluna_seq else "",
                "modelo": _texto_celula(linha.get("MODELO")),
                "posicao": posicao + 1,
            })
        return veiculos

    def coluna_sequencia(self) -> str | None:
        """Acha a coluna de sequencia de producao, qualquer que seja o nome."""
        if self._coluna_seq is not False:
            return self._coluna_seq

        self._coluna_seq = None
        normalizadas = {
            str(c).strip().upper().replace(".", "").replace("_", "").replace(" ", ""): c
            for c in self.df.columns
        }
        for candidata in COLUNAS_SEQUENCIA:
            if candidata in normalizadas:
                self._coluna_seq = normalizadas[candidata]
                break
        return self._coluna_seq


# ---------------------------------------------------------------------------
# 5. Estruturas de resultado (o que a interface em colunas consome)
# ---------------------------------------------------------------------------
@dataclass
class ResultadoTarefa:
    nome: str
    contagem: int
    tempo: float
    percentual: float
    itens: list = field(default_factory=list)   # POPID/SEQ dos veiculos contados

    @property
    def percentual_fmt(self) -> str:
        return f"{self.percentual:.2f}%"


@dataclass
class ResultadoVariante:
    id: str
    nome: str
    tarefas: list[ResultadoTarefa] = field(default_factory=list)
    absenteismo: ResultadoTarefa | None = None

    @property
    def subtotal(self) -> float:
        return sum(t.percentual for t in self.tarefas)

    @property
    def total(self) -> float:
        extra = self.absenteismo.percentual if self.absenteismo else 0.0
        return self.subtotal + extra

    @property
    def total_fmt(self) -> str:
        return f"{self.total:.2f}%"

    def linhas(self) -> list[ResultadoTarefa]:
        """Tarefas + absenteismo, na ordem em que devem aparecer na coluna."""
        return self.tarefas + ([self.absenteismo] if self.absenteismo else [])


@dataclass
class ResultadoAnalise:
    popid: str
    primeira_linha: int
    ultima_linha: int
    total_linhas: int          # linhas efetivamente analisadas
    producao_informada: int
    linhas_originais: int
    mix: dict[str, int] = field(default_factory=dict)
    rotulos_mix: dict[str, str] = field(default_factory=dict)
    itens_mix: dict[str, list] = field(default_factory=dict)
    tem_coluna_seq: bool = False
    demanda: list = field(default_factory=list)   # carga (%) exigida por veículo
    veiculos: list = field(default_factory=list)  # POPID/SEQ de cada posição
    faltas: int = 0                 # absenteísmo do setor (pessoas ausentes)
    postos_ativos: int = 0          # variantes que seguem operando
    postos_originais: int = 0
    alertas: list = field(default_factory=list)
    variantes: dict[str, ResultadoVariante] = field(default_factory=dict)
    avisos: list[str] = field(default_factory=list)
    catalogo: tuple = ()   # catalogo usado nesta analise

    # -- compatibilidade com o codigo que ja existe -------------------------
    def to_legacy_dict(self) -> dict:
        legado: dict = {
            "LINHAS_ORIGINAIS": self.linhas_originais,
            "PRIMEIRA_LINHA": self.primeira_linha,
            "ULTIMA_LINHA": self.ultima_linha,
            "TOTAL_LINHAS": self.total_linhas,
            "POPID": self.popid,
        }
        legado.update(self.mix)
        for variante in (self.catalogo or CATALOGO_PADRAO):
            resultado = self.variantes.get(variante.id)
            if resultado is None:
                continue
            for tarefa, calculada in zip(variante.tarefas, resultado.tarefas):
                chave = ALIASES_LEGADO.get(tarefa.chave_legado, tarefa.chave_legado)
                legado[chave] = (calculada.contagem, calculada.percentual_fmt)
        return legado


# ---------------------------------------------------------------------------
# 6. Analise
# ---------------------------------------------------------------------------
def _normalizar_popid(valor) -> str:
    if isinstance(valor, float) and float(valor).is_integer():
        return str(int(valor))
    return str(valor).strip()


def _localizar_popid(df: pd.DataFrame, popid) -> tuple[str, int]:
    """Devolve (popid_normalizado, indice_posicional_da_primeira_ocorrencia)."""
    if "POPID" not in df.columns:
        raise ValueError("Coluna 'POPID' nao encontrada nos dados")

    serie = df["POPID"].map(_normalizar_popid)
    alvo = _normalizar_popid(popid)

    ocorrencias = serie.eq(alvo)
    if not ocorrencias.any():
        alvo_sem_zeros = alvo.lstrip("0")
        ocorrencias = serie.str.lstrip("0").eq(alvo_sem_zeros)
        if not ocorrencias.any():
            raise ValueError(f"POPID {alvo} nao encontrado")
        alvo = alvo_sem_zeros

    posicao = int(ocorrencias.to_numpy().argmax())
    return alvo, posicao


def analisar_mix(
    df: pd.DataFrame,
    popid,
    producao,
    takt: float = SEGUNDOS_DISPONIVEIS_POR_UNIDADE,
    absenteismo: dict[str, float] | None = None,
    base_percentual: str = "informada",
    catalogo: Sequence[Variante] | None = None,
    contador_mona: str | None = None,
    setor: str | None = None,
) -> ResultadoAnalise:
    """
    Analisa `producao` linhas a partir da primeira ocorrencia do POPID.

    absenteismo    : {"V1": 2, "V2": 0, ...} numero de faltas por variante.
    base_percentual: "informada" usa a producao digitada no denominador
                     (comportamento historico); "analisada" usa o numero real
                     de linhas quando o intervalo e truncado no fim do arquivo.
    catalogo       : lista de Variante. Por padrao usa `catalogo_ativo(setor)`,
                     que le o catalogo.json (por setor) editavel pela interface
                     de Tarefas — assim cada setor mostra as tarefas dele, nao
                     as da FA2.
    contador_mona  : texto avulso usado no cruzamento com a MONA nesta analise
                     (o "Outro" do menu, quando digitado). Quando informado e
                     ainda nao for uma tag cadastrada em tags_mona.json, um
                     contador extra com esse texto entra no CPM MIX so desta
                     vez. As tags fixas (editaveis em /tags-mona) entram
                     sempre, sem precisar digitar nada.
    setor          : id do setor (ex.: "fa2", "1-2") — decide qual catalogo
                     de VARIANTES/TAREFAS (`catalogo_ativo`), qual catalogo de
                     itens do CPM MIX (`catalogo_mix_do_setor`) e quais tags
                     da MONA (`tags_mona_ativas`) entram na analise. Sem isso,
                     cai no padrao da FA2, igual sempre foi.
    """
    catalogo = tuple(catalogo) if catalogo else catalogo_ativo(setor)
    if df is None or getattr(df, "empty", True):
        raise ValueError("DataFrame vazio ou invalido")

    try:
        producao = int(producao)
    except (TypeError, ValueError):
        raise ValueError("O parametro 'producao' deve ser um numero inteiro")
    if producao <= 0:
        raise ValueError("O parametro 'producao' deve ser maior que zero")

    avisos: list[str] = []
    popid_str, inicio = _localizar_popid(df, popid)
    fim = min(inicio + producao, len(df))
    recorte = df.iloc[inicio:fim].copy()

    linhas_analisadas = len(recorte)
    if linhas_analisadas < producao:
        avisos.append(
            f"Intervalo truncado: {linhas_analisadas} linhas disponiveis de "
            f"{producao} solicitadas."
        )

    base = producao if base_percentual == "informada" else linhas_analisadas

    ctx = ContextoDados(recorte)

    tags_ativas = tags_mona_ativas(setor)
    catalogo_mix = catalogo_mix_do_setor(setor) + tuple(
        _contador_texto_opcionais(tag.texto) for tag in tags_ativas
    )
    textos_fixos = {tag.texto for tag in tags_ativas}
    if contador_mona and contador_mona not in textos_fixos:
        catalogo_mix = catalogo_mix + (_contador_texto_opcionais(contador_mona),)

    itens_mix = {c.chave: ctx.identificar(c.criterio) for c in catalogo_mix}
    mix = {chave: len(itens) for chave, itens in itens_mix.items()}
    rotulos = {c.chave: c.rotulo for c in catalogo_mix}

    faltas = absenteismo or {}
    # Tempo exigido por veículo: quanto de trabalho cada carro traz ao passar
    # pelo setor. E o que desenha a curva de demanda da linha do tempo.
    tempo_por_posicao = [0.0] * linhas_analisadas

    variantes: dict[str, ResultadoVariante] = {}
    for variante in catalogo:
        calculadas = []
        for tarefa in variante.tarefas:
            itens = ctx.identificar(tarefa.criterio)
            contagem = len(itens)
            for item in itens:
                indice = item["posicao"] - 1
                if 0 <= indice < linhas_analisadas:
                    tempo_por_posicao[indice] += tarefa.tempo
            calculadas.append(ResultadoTarefa(
                nome=tarefa.nome,
                contagem=contagem,
                tempo=tarefa.tempo,
                percentual=carga_percentual(contagem, tarefa.tempo, base, takt),
                itens=itens,
            ))

        n_faltas = float(faltas.get(variante.id, 0) or 0)
        resultado_abs = None
        if n_faltas:
            resultado_abs = ResultadoTarefa(
                nome="Absenteismo",
                contagem=int(n_faltas) if float(n_faltas).is_integer() else n_faltas,
                tempo=TEMPO_ABSENTEISMO,
                percentual=carga_percentual(n_faltas, TEMPO_ABSENTEISMO, base, takt),
            )

        variantes[variante.id] = ResultadoVariante(
            id=variante.id,
            nome=variante.nome,
            tarefas=calculadas,
            absenteismo=resultado_abs,
        )

    if ctx.colunas_ausentes:
        avisos.append(
            "Colunas ausentes na planilha (contadas como zero): "
            + ", ".join(sorted(ctx.colunas_ausentes))
        )

    return ResultadoAnalise(
        popid=popid_str,
        primeira_linha=inicio + 1,
        ultima_linha=fim,
        total_linhas=linhas_analisadas,
        producao_informada=producao,
        linhas_originais=len(df),
        demanda=[carga_percentual(1, tempo, 1, takt) for tempo in tempo_por_posicao],
        veiculos=ctx.identificar_todos(),
        mix=mix,
        rotulos_mix=rotulos,
        itens_mix=itens_mix,
        tem_coluna_seq=ctx.coluna_sequencia() is not None,
        variantes=variantes,
        avisos=avisos,
        catalogo=catalogo,
    )


def analisar_88_linhas_completas(df, popid, producao, caminho_arquivo_origem=None,
                                 sheet_name=0):
    """Wrapper compativel com o codigo antigo: devolve o dicionario de sempre."""
    try:
        if caminho_arquivo_origem:
            df_visivel = carregar_excel_sem_linhas_ocultas(caminho_arquivo_origem, sheet_name)
            if df_visivel is None:
                return {"ERRO": "Falha ao filtrar linhas ocultas"}
            df = df_visivel
        return analisar_mix(df, popid, producao).to_legacy_dict()
    except ValueError as erro:
        return {"ERRO": str(erro)}
    except Exception as erro:  # noqa: BLE001 - contrato antigo
        return {"ERRO": f"Erro durante analise: {erro}"}


# ---------------------------------------------------------------------------
# 7. Leitura de planilhas
# ---------------------------------------------------------------------------
def _normalizar_cabecalho(bruto: Sequence, origem: str = "") -> tuple[list[str], list[str]]:
    """
    Deixa os nomes de coluna unicos e preenchidos.

    Cabecalho de planilha real costuma ter celula vazia e nome repetido. Com
    nome repetido o pandas nao consegue alinhar dois DataFrames e o concat
    quebra com "Reindexing only valid with uniquely valued Index objects".
    Aqui a coluna vazia vira COLUNA_<posicao> e a repetida ganha sufixo
    (.1, .2), que e o mesmo que pd.read_excel faz.
    """
    nomes: list[str] = []
    avisos: list[str] = []
    vistos: dict[str, int] = {}
    vazias = 0
    repetidas: list[str] = []

    for posicao, valor in enumerate(bruto, 1):
        nome = "" if valor is None else str(valor).strip()
        if not nome or nome.lower() == "none":
            nome = f"COLUNA_{posicao}"
            vazias += 1

        if nome in vistos:
            vistos[nome] += 1
            repetidas.append(nome)
            nome = f"{nome}.{vistos[nome]}"
        else:
            vistos[nome] = 0

        nomes.append(nome)

    prefixo = f"{origem}: " if origem else ""
    if vazias:
        avisos.append(f"{prefixo}{vazias} coluna(s) sem nome renomeada(s) para COLUNA_<posição>.")
    if repetidas:
        unicas = sorted(set(repetidas))
        avisos.append(
            f"{prefixo}nome(s) de coluna repetido(s) — {', '.join(unicas)} — "
            f"receberam sufixo (.1, .2). Confira se a planilha está correta."
        )
    return nomes, avisos


def _alinhar_quadros(quadros: Sequence[pd.DataFrame]) -> tuple[list[pd.DataFrame], list[str]]:
    """
    Deixa todos os DataFrames com as mesmas colunas, na mesma ordem.

    Usa a UNIAO das colunas: o que existe so em um arquivo vira vazio nos
    outros, em vez de ser descartado silenciosamente.
    """
    avisos: list[str] = []
    ordem: list[str] = []
    for quadro in quadros:
        for coluna in quadro.columns:
            if coluna not in ordem:
                ordem.append(coluna)

    for i, quadro in enumerate(quadros, 1):
        faltando = [c for c in ordem if c not in quadro.columns]
        if faltando:
            avisos.append(
                f"Planilha {i} não tem a(s) coluna(s) {', '.join(faltando)} — "
                f"preenchida(s) como vazio."
            )

    return [q.reindex(columns=ordem) for q in quadros], avisos


def carregar_excel_sem_linhas_ocultas(
    caminho_arquivo,
    sheet_name=0,
    linhas_ignorar: Iterable[int] | None = None,
    ignorar_ocultas: bool = True,
    coluna_filtro: str = "ENTRADA LINHA REAL",
    valores_excluidos: Sequence[str] = ("TVL",),
):
    """
    Le a planilha ignorando linhas ocultas e linhas cuja coluna `coluna_filtro`
    contenha algum dos `valores_excluidos`.

    Correcoes em relacao a versao anterior:
      * data_only=True  -> celulas com formula devolvem o valor, nao "=A1*2"
      * o cabecalho e sempre capturado, mesmo que a linha 1 esteja oculta
      * o workbook e fechado ao final
    """
    workbook = None
    try:
        workbook = load_workbook(caminho_arquivo, data_only=True)
        sheet = workbook.worksheets[sheet_name] if isinstance(sheet_name, int) else workbook[sheet_name]

        ignorar = set(linhas_ignorar or ())
        excluidos = tuple(v.upper() for v in valores_excluidos)

        cabecalho, avisos_cabecalho = _normalizar_cabecalho(
            [c.value for c in sheet[1]], os.path.basename(str(caminho_arquivo))
        )
        try:
            idx_filtro = cabecalho.index(coluna_filtro)
        except ValueError:
            idx_filtro = None

        if not any(v is not None for v in cabecalho):
            return pd.DataFrame()

        dados = []
        for row_idx, row in enumerate(sheet.iter_rows(values_only=True), 1):
            if row_idx == 1:
                continue
            if ignorar_ocultas and sheet.row_dimensions[row_idx].hidden:
                continue
            if (row_idx - 1) in ignorar:
                continue
            if idx_filtro is not None and len(row) > idx_filtro:
                valor = row[idx_filtro]
                if isinstance(valor, str) and any(e in valor.upper() for e in excluidos):
                    continue
            dados.append(row)

        df = pd.DataFrame(dados, columns=cabecalho)
        df.attrs["avisos"] = avisos_cabecalho
        return df
    except Exception as erro:  # noqa: BLE001
        logger.warning("erro ao carregar %s: %s", caminho_arquivo, erro)
        return None
    finally:
        if workbook is not None:
            workbook.close()


def concatenar_planilhas_em_memoria(arquivo1, arquivo2, verbose: bool = True):
    """Concatena duas planilhas alinhando as colunas pela uniao."""
    df1 = carregar_excel_sem_linhas_ocultas(arquivo1)
    df2 = carregar_excel_sem_linhas_ocultas(arquivo2, linhas_ignorar=[1])

    if df1 is None or df2 is None:
        return None

    if verbose:
        logger.debug("linhas carregadas do arquivo 1: %s", len(df1))
        logger.debug("linhas carregadas do arquivo 2: %s", len(df2))

    avisos = list(df1.attrs.get("avisos", [])) + list(df2.attrs.get("avisos", []))
    quadros = [_limpar_colunas(df1, "Planilha 1"), _limpar_colunas(df2, "Planilha 2")]
    quadros, avisos_alinhamento = _alinhar_quadros(quadros)
    avisos += avisos_alinhamento

    resultado = pd.concat(quadros, ignore_index=True)
    resultado.attrs["avisos"] = avisos
    if verbose:
        for aviso in avisos:
            logger.info("aviso: %s", aviso)
    return resultado


def _limpar_colunas(df: pd.DataFrame, nome: str = "planilha") -> pd.DataFrame:
    """Padroniza nomes, remove colunas/indices duplicados e reindexa."""
    df = df.copy()
    df.columns = [str(c).strip() for c in df.columns]

    # Rede de seguranca: se ainda restar nome repetido (ex.: DataFrame vindo
    # de outra origem), renomeia em vez de descartar a coluna.
    if pd.Index(df.columns).has_duplicates:
        df.columns, _ = _normalizar_cabecalho(df.columns, nome)

    if df.index.duplicated().any():
        df = df[~df.index.duplicated(keep="first")]

    return df.reset_index(drop=True)


# Data no nome do arquivo: "mix a partir do pop-id 775301 ( 10_09 )".
# Aceita ( 10_09 ), (10-09), (10/09/2026) e variacoes com ou sem espaco.
_DATA_NO_NOME = re.compile(
    r"[(\[]\s*(\d{1,2})\s*[_\-./]\s*(\d{1,2})(?:\s*[_\-./]\s*(\d{2,4}))?\s*[)\]]"
)


def data_do_nome(caminho: str) -> tuple[int, int, int] | None:
    """
    Extrai (ano, mes, dia) do nome do arquivo. Ano ausente vira 0 — a ordem
    entre dois arquivos do mesmo periodo nao depende dele.
    """
    achado = _DATA_NO_NOME.search(os.path.basename(str(caminho)))
    if not achado:
        return None

    dia, mes, ano = achado.group(1), achado.group(2), achado.group(3)
    dia, mes = int(dia), int(mes)
    if not (1 <= dia <= 31 and 1 <= mes <= 12):
        return None

    if ano:
        ano = int(ano)
        if ano < 100:
            ano += 2000
    else:
        ano = 0
    return (ano, mes, dia)


def ordenar_por_data(caminhos: Sequence[str]) -> tuple[list[str], list[str]]:
    """
    Coloca as planilhas em ordem cronologica pela data do nome — a mais antiga
    primeiro, que e o que a concatenacao espera.

    O navegador entrega os arquivos na ordem em que o sistema listou, nao na
    ordem em que interessa; por isso a ordem vinha errada quando o envio era
    manual.
    """
    avisos: list[str] = []
    datados = [(data_do_nome(c), c) for c in caminhos]
    sem_data = [c for data, c in datados if data is None]

    if sem_data:
        avisos.append(
            "Sem data no nome (esperado algo como \"( 10_09 )\"): "
            + ", ".join(os.path.basename(c) for c in sem_data)
            + ". A ordem de envio foi mantida."
        )
        return list(caminhos), avisos

    # So dia e mes: um par 31/12 e 01/01 e virada de ano, nao meio ano de
    # distancia. Quando os dois extremos estao longe demais, o mais tarde no
    # calendario e na verdade do ano anterior.
    if all(data[0] == 0 for data, _ in datados):
        dias = [mes * 31 + dia for (_, mes, dia), _ in datados]
        if max(dias) - min(dias) > 182:
            corte = (max(dias) + min(dias)) / 2
            datados = [((-1 if (mes * 31 + dia) > corte else 0, mes, dia), c)
                       for (_, mes, dia), c in datados]

    ordenados = [c for _, c in sorted(datados, key=lambda item: item[0])]

    if ordenados != list(caminhos):
        avisos.append(
            "Planilhas reordenadas pela data do nome: "
            + " → ".join(os.path.basename(c) for c in ordenados)
        )
    return ordenados, avisos


def concatenar_planilhas(caminhos: Sequence[str], verbose: bool = False) -> pd.DataFrame:
    """
    Concatena N planilhas (ordem cronologica: mais antiga primeiro), unindo o
    conjunto de colunas. E o caminho usado pela captura automatica do Outlook,
    que pode trazer 1, 2 ou 3 arquivos.
    """
    caminhos = [c for c in caminhos if c]
    if not caminhos:
        raise ValueError("Nenhuma planilha para concatenar")

    quadros = []
    avisos: list[str] = []
    for i, caminho in enumerate(caminhos, 1):
        df = carregar_excel_sem_linhas_ocultas(caminho)
        if df is None:
            raise ValueError(f"Falha ao ler a planilha {i}: {os.path.basename(caminho)}")
        avisos += list(df.attrs.get("avisos", []))
        quadros.append(_limpar_colunas(df, f"Planilha {i}"))

    quadros, avisos_alinhamento = _alinhar_quadros(quadros)
    avisos += avisos_alinhamento

    resultado = pd.concat(quadros, ignore_index=True, sort=False)
    resultado.attrs["avisos"] = avisos
    if verbose:
        logger.debug("concatenadas %s planilhas -> %s", len(quadros), resultado.shape)
        for aviso in avisos:
            logger.info("aviso: %s", aviso)
    return resultado


def gerar_arquivo_concatenado(arquivo1, arquivo2, pasta_destino):
    """Gera o xlsx concatenado preservando estilos. Devolve (ok, caminho_ou_erro)."""
    if not all(os.path.isfile(f) for f in (arquivo1, arquivo2)):
        return False, "Um ou ambos os arquivos de entrada nao existem"

    wb1 = wb2 = None
    try:
        Path(pasta_destino).mkdir(parents=True, exist_ok=True)

        wb1 = load_workbook(arquivo1)
        wb2 = load_workbook(arquivo2)
        sheet1, sheet2 = wb1.active, wb2.active

        wb_saida = Workbook()
        sheet_saida = wb_saida.active
        sheet_saida.title = "Dados Concatenados"

        def copiar(origem, destino):
            destino.value = origem.value
            if origem.has_style:
                destino.font = copy(origem.font)
                destino.border = copy(origem.border)
                destino.fill = copy(origem.fill)
                destino.number_format = origem.number_format
                destino.protection = copy(origem.protection)
                destino.alignment = copy(origem.alignment)

        for row_idx, row in enumerate(sheet1.iter_rows(), 1):
            for col_idx, cell in enumerate(row, 1):
                copiar(cell, sheet_saida.cell(row=row_idx, column=col_idx))

        offset = sheet_saida.max_row
        for row_idx, row in enumerate(sheet2.iter_rows(min_row=3), offset + 1):
            for col_idx, cell in enumerate(row, 1):
                copiar(cell, sheet_saida.cell(row=row_idx, column=col_idx))

        for col in sheet1.columns:
            letra = col[0].column_letter
            largura = sheet1.column_dimensions[letra].width
            if largura:
                sheet_saida.column_dimensions[letra].width = largura

        nome = f"Planilha_Concatenada_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        caminho = os.path.join(pasta_destino, nome)
        wb_saida.save(caminho)
        return True, caminho
    except PermissionError:
        return False, "Arquivo de saida aberto em outro programa. Feche-o e tente novamente."
    except Exception as erro:  # noqa: BLE001
        return False, f"Erro ao gerar arquivo concatenado: {erro}"
    finally:
        for wb in (wb1, wb2):
            if wb is not None:
                wb.close()


# ---------------------------------------------------------------------------
# 8. Relatorio em texto (mantido para quem ainda consome string)
# ---------------------------------------------------------------------------
def _num(valor: float, casas: int = 2) -> str:
    """Numero com virgula decimal, para colar no Excel em portugues."""
    return f"{valor:.{casas}f}".replace(".", ",")


def exportar_tsv(resultado: ResultadoAnalise) -> str:
    """
    Resultado em colunas separadas por TAB. Colando no Excel, cada campo cai
    em uma celula sem precisar de "texto para colunas".
    """
    linhas = [
        f"POPID\t{resultado.popid}",
        f"Linhas analisadas\t{resultado.primeira_linha} a {resultado.ultima_linha}",
        f"Veiculos\t{resultado.total_linhas}",
        f"Absenteismo\t{resultado.faltas}",
        f"Variantes ativas\t{resultado.postos_ativos} de {resultado.postos_originais}",
    ] + [f"ALERTA\t{a['texto']} - {a['detalhe']}" for a in resultado.alertas] + [
        "",
        "Variante\tTarefa\tQtde\tTempo (s)\tCarga (%)",
    ]
    for variante in resultado.variantes.values():
        for tarefa in variante.linhas():
            linhas.append("\t".join([
                variante.nome, tarefa.nome, str(tarefa.contagem),
                _num(tarefa.tempo, 0), _num(tarefa.percentual),
            ]))
        linhas.append("\t".join([variante.nome, "TOTAL", "", "",
                                  _num(variante.total)]))

    linhas += ["", "CPM MIX", "Item\tQuantidade"]
    for chave, valor in resultado.mix.items():
        linhas.append(f"{resultado.rotulos_mix.get(chave, chave)}\t{valor}")

    return "\n".join(linhas)


def exportar_texto(resultado: ResultadoAnalise) -> str:
    """Bloco legivel para colar em e-mail, Teams ou WhatsApp."""
    linhas = [
        "ANALISE DE MIX",
        f"POPID {resultado.popid} | linhas {resultado.primeira_linha}-"
        f"{resultado.ultima_linha} | {resultado.total_linhas} veiculos",
    ]
    if resultado.faltas:
        linhas.append(
            f"Absenteismo: {resultado.faltas} pessoa(s) | "
            f"{resultado.postos_ativos} de {resultado.postos_originais} variantes"
        )
    for alerta in resultado.alertas:
        linhas.append(f"** {alerta['texto'].upper()} - {alerta['detalhe']}")
    linhas.append("")
    for variante in resultado.variantes.values():
        linhas.append(f"{variante.nome.upper()} - {_num(variante.total)}%")
        for tarefa in variante.linhas():
            linhas.append(
                f"  {tarefa.nome:.<32} {_num(tarefa.percentual):>7}%  ({tarefa.contagem}x)"
            )
        linhas.append("")

    linhas.append("CPM MIX")
    itens = [(resultado.rotulos_mix.get(c, c), v)
             for c, v in resultado.mix.items() if v]
    largura = max((len(rotulo) for rotulo, _ in itens), default=0) + 2
    for rotulo, valor in itens:
        linhas.append(f"  {rotulo:.<{largura}} {valor}")

    for aviso in resultado.avisos:
        linhas.append(f"[aviso] {aviso}")

    return "\n".join(linhas)


def suavizar(valores: Sequence[float], janela: int = 5) -> list[float]:
    """Media movel centrada — a curva crua e serrilhada demais para ler."""
    if janela < 2 or not valores:
        return list(valores)
    metade = janela // 2
    suave = []
    for i in range(len(valores)):
        inicio = max(0, i - metade)
        fim = min(len(valores), i + metade + 1)
        trecho = valores[inicio:fim]
        suave.append(sum(trecho) / len(trecho))
    return suave


def _segundos_do_relogio(hhmm: str, padrao: int = 0) -> int:
    try:
        hora, minuto = (int(p) for p in str(hhmm).split(":")[:2])
        return hora * 3600 + minuto * 60
    except (ValueError, AttributeError):
        return padrao


def duracao_do_turno(inicio: str = INICIO_TURNO_PADRAO,
                     fim: str = FIM_TURNO_PADRAO) -> int:
    """Duração bruta do turno em segundos. Turno que vira a noite é tratado."""
    inicio_s = _segundos_do_relogio(inicio, 7 * 3600 + 5 * 60)
    fim_s = _segundos_do_relogio(fim, 16 * 3600 + 5 * 60)
    if fim_s <= inicio_s:
        fim_s += 24 * 3600
    return fim_s - inicio_s


def _pausas_em_segundos(inicio: str, fim: str, pausas=PAUSAS_PADRAO
                        ) -> list[tuple[int, int]]:
    """Pausas dentro do turno, em segundos desde a meia-noite do dia inicial."""
    inicio_s = _segundos_do_relogio(inicio, 7 * 3600 + 5 * 60)
    limite = inicio_s + duracao_do_turno(inicio, fim)

    intervalos = []
    for pausa_inicio, pausa_fim in pausas or ():
        p1 = _segundos_do_relogio(pausa_inicio)
        p2 = _segundos_do_relogio(pausa_fim)
        if p1 < inicio_s:          # turno que vira a noite
            p1 += 24 * 3600
            p2 += 24 * 3600
        if p2 <= p1:
            p2 += 24 * 3600
        # so o que cai dentro do turno conta
        p1, p2 = max(p1, inicio_s), min(p2, limite)
        if p2 > p1:
            intervalos.append((p1, p2))

    return sorted(intervalos)


def tempo_produtivo(inicio: str = INICIO_TURNO_PADRAO,
                    fim: str = FIM_TURNO_PADRAO, pausas=PAUSAS_PADRAO) -> int:
    """Segundos de turno descontadas as pausas."""
    parado = sum(p2 - p1 for p1, p2 in _pausas_em_segundos(inicio, fim, pausas))
    return max(duracao_do_turno(inicio, fim) - parado, 1)


def ritmo_do_turno(total_veiculos: int, inicio: str = INICIO_TURNO_PADRAO,
                   fim: str = FIM_TURNO_PADRAO, pausas=PAUSAS_PADRAO) -> float:
    """
    Segundos por veículo derivados do turno, ja sem as pausas.

    Nao e o mesmo que SEGUNDOS_DISPONIVEIS_POR_UNIDADE: aquele e o tempo de
    ciclo usado na formula de carga (fixo). Este e so o ritmo do relogio, para
    posicionar os veiculos no eixo de horas.
    """
    if total_veiculos <= 0:
        return float(SEGUNDOS_DISPONIVEIS_POR_UNIDADE)
    return tempo_produtivo(inicio, fim, pausas) / total_veiculos


def relogio_apos(produtivo: float, inicio: str = INICIO_TURNO_PADRAO,
                 fim: str = FIM_TURNO_PADRAO, pausas=PAUSAS_PADRAO) -> float:
    """
    Converte tempo produtivo decorrido em horario de relogio.

    Cada pausa ja ultrapassada empurra o relogio para frente — e por isso que
    o ultimo veiculo sai perto do fim do turno mesmo com as paradas.
    """
    momento = _segundos_do_relogio(inicio, 7 * 3600 + 5 * 60) + produtivo
    for p1, p2 in _pausas_em_segundos(inicio, fim, pausas):
        if momento >= p1:
            momento += p2 - p1
    return momento


def _hhmm(segundos: float) -> str:
    segundos = int(segundos) % (24 * 3600)
    return f"{segundos // 3600:02d}:{segundos % 3600 // 60:02d}"


def horario_da_posicao(posicao: int, inicio_turno: str = INICIO_TURNO_PADRAO,
                       segundos_por_veiculo: float = SEGUNDOS_DISPONIVEIS_POR_UNIDADE,
                       fim_turno: str = FIM_TURNO_PADRAO,
                       pausas=PAUSAS_PADRAO) -> str:
    """Horário estimado em que o veículo da posição passa pelo setor."""
    return _hhmm(relogio_apos(posicao * segundos_por_veiculo,
                              inicio_turno, fim_turno, pausas))


# ---------------------------------------------------------------------------
# Absenteismo do setor: redistribuicao das tarefas
# ---------------------------------------------------------------------------
# Quantos postos seguem operando conforme o numero de ausentes. E uma tabela,
# nao uma subtracao: com 1 falta a distribuicao continua normal.
POSTOS_POR_FALTAS = {0: None, 1: None, 2: 2}   # None = todos os postos
POSTOS_MINIMO = 1                              # 3 ou mais faltas

# A partir de quantos gargalos se considera "varios" (regra de 3+ faltas).
GARGALOS_VARIOS = 2

MSG_RECURSO = "Necessário recurso"
MSG_DESCOBERTA = "Linha descoberta para várias tarefas"


def postos_ativos_por_faltas(faltas: int, total_postos: int,
                             tabela: dict | None = None) -> int:
    """Quantas variantes continuam operando com `faltas` pessoas ausentes."""
    if faltas <= 0:
        return total_postos
    tabela = POSTOS_POR_FALTAS if tabela is None else tabela
    escolha = tabela.get(faltas, POSTOS_MINIMO)
    if escolha is None:
        return total_postos
    return max(POSTOS_MINIMO, min(escolha, total_postos))


def _distribuir(tarefas: list[ResultadoTarefa],
                destinos: list[ResultadoVariante]) -> None:
    """
    Reparte as tarefas entre os postos que sobraram, sempre jogando a proxima
    no posto mais leve. Nao e a divisao otima (isso e NP-dificil), mas e a
    heuristica usual e deixa os totais proximos entre si.
    """
    for destino in destinos:
        destino.tarefas = []

    for tarefa in sorted(tarefas, key=lambda t: t.percentual, reverse=True):
        mais_leve = min(destinos, key=lambda d: d.subtotal)
        mais_leve.tarefas.append(tarefa)

    for destino in destinos:
        destino.tarefas.sort(key=lambda t: t.percentual, reverse=True)


def aplicar_absenteismo(resultado: ResultadoAnalise, faltas: int,
                        n_gargalos: int = 0, tabela_postos: dict | None = None,
                        gargalos_varios: int | None = None) -> ResultadoAnalise:
    """
    Aplica o absenteismo do setor: reduz os postos ativos, redistribui as
    tarefas entre os que sobraram e monta os alertas.

    Diferente do modelo antigo, o absenteismo nao entra mais como um tempo
    fixo somado a variante — quem cobre a ausencia sao os postos restantes, e
    e isso que a redistribuicao mostra.
    """
    faltas = max(0, int(faltas or 0))
    variantes = list(resultado.variantes.values())

    resultado.faltas = faltas
    resultado.postos_originais = len(variantes)
    resultado.postos_ativos = postos_ativos_por_faltas(
        faltas, len(variantes), tabela_postos)
    resultado.alertas = []

    if not variantes:
        return resultado

    if resultado.postos_ativos < len(variantes):
        ativos = variantes[:resultado.postos_ativos]
        desativados = variantes[resultado.postos_ativos:]

        todas = []
        for variante in variantes:
            for tarefa in variante.tarefas:
                tarefa.origem = variante.id
                todas.append(tarefa)

        _distribuir(todas, ativos)
        for variante in desativados:
            variante.tarefas = []
        resultado.variantes = {v.id: v for v in ativos}

    for variante in resultado.variantes.values():
        variante.absenteismo = None

    # --- alertas ---
    if faltas >= 3:
        resultado.alertas.append({
            "nivel": "alto", "texto": MSG_DESCOBERTA,
            "detalhe": f"{faltas} ausências — todas as tarefas em um único posto.",
        })
        if n_gargalos >= (GARGALOS_VARIOS if gargalos_varios is None
                          else gargalos_varios):
            resultado.alertas.append({
                "nivel": "critico", "texto": MSG_RECURSO,
                "detalhe": f"{faltas} ausências e {n_gargalos} momentos de gargalo.",
            })
    elif faltas == 2 and n_gargalos >= 1:
        resultado.alertas.append({
            "nivel": "critico", "texto": MSG_RECURSO,
            "detalhe": f"2 ausências e {n_gargalos} momento(s) de gargalo.",
        })

    return resultado


def formatar_relatorio_texto(resultado: ResultadoAnalise) -> str:
    linhas = [
        "=== RELATORIO DE ANALISE DO MIX ===",
        f"POPID de referencia: {resultado.popid}",
        f"Intervalo analisado: linhas {resultado.primeira_linha} a {resultado.ultima_linha}",
        f"Total de linhas analisadas: {resultado.total_linhas}",
        "",
        "=== CPM MIX ===",
    ]
    for chave, valor in resultado.mix.items():
        linhas.append(f"{resultado.rotulos_mix.get(chave, chave)} - {valor}")

    linhas += ["", "=== WORKLOAD VARIANTES ==="]
    for variante in resultado.variantes.values():
        linhas.append(f"{variante.nome} - {variante.total_fmt}")

    for aviso in resultado.avisos:
        linhas.append(f"[aviso] {aviso}")

    return "\n".join(linhas)
