"""Fixtures compartilhadas pelos testes de cruzamento com a MONA."""
from openpyxl import Workbook


def escrever_mona_pos(caminho, linhas):
    """
    Grava uma planilha MONA no formato tecnico da Funcao 4:
      linha 1 = cabecalhos tecnicos (POPID, POS_10, POS_20...)
      linha 2 = nome do item de cada coluna POS_
      linha 3+ = dados

    `linhas` e uma lista de dicts {"POPID": ..., "<nome do item>": valor}.
    As colunas POS_ sao geradas automaticamente a partir dos nomes de item
    encontrados em `linhas`.
    """
    nomes_item = []
    for linha in linhas:
        for chave in linha:
            if chave != "POPID" and chave not in nomes_item:
                nomes_item.append(chave)

    wb = Workbook()
    ws = wb.active
    cabecalho_tecnico = ["POPID"] + [f"POS_{10 * (i + 1)}" for i in range(len(nomes_item))]
    ws.append(cabecalho_tecnico)
    ws.append([""] + nomes_item)
    for linha in linhas:
        ws.append([linha.get("POPID")] + [linha.get(nome) for nome in nomes_item])
    wb.save(caminho)


def escrever_mona_lista_simples(caminho, popids, coluna_extra=None):
    """Grava uma planilha MONA "fora do padrao": so uma coluna POPID comum."""
    wb = Workbook()
    ws = wb.active
    cabecalho = ["POPID"] + ([coluna_extra] if coluna_extra else [])
    ws.append(cabecalho)
    for popid in popids:
        ws.append([popid] + (["x"] if coluna_extra else []))
    wb.save(caminho)


def escrever_mix_fa2(caminho, linhas):
    """
    Grava um MIX minimo no formato que o Filtro MIX (FA2) espera.
    `linhas` e uma lista de dicts com as colunas do MIX.
    """
    colunas = ["SEQ.", "REPLAN", "POPID", "CHASSI", "MODELO",
              "SUSP TRASEIRA", "COMPRIM", "OPCIONAIS", "MU", "RPX",
              "BARRA 2°E DIANT", "BARRA TRASEIRA"]
    wb = Workbook()
    ws = wb.active
    ws.append(colunas)
    for i, linha in enumerate(linhas, start=1):
        ws.append([
            i, "", linha.get("POPID"), linha.get("CHASSI", "111"),
            linha.get("MODELO", "MOD-A"), linha.get("SUSP TRASEIRA", "T"),
            linha.get("COMPRIM", 7000), linha.get("OPCIONAIS", ""),
            linha.get("MU", "BR01"), linha.get("RPX", ""), "", "",
        ])
    wb.save(caminho)
