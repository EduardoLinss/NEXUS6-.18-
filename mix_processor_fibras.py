# -*- coding: utf-8 -*-
"""Geração do MIX DE FIBRAS a partir do MIX geral universal."""

from datetime import date

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


FONTE_PADRAO = "Arial"
FONTE_CODIGO_BARRAS = "Free 3 of 9 Extended"

CINZA_CABECALHO = "FFC0C0C0"
AZUL = "FF0070C0"
VERDE = "FF00B050"
VERDE_CLARO = "FFCCFFCC"
AMARELO = "FFFFFF00"

CENTRO = Alignment(horizontal="center", vertical="center")
BORDA = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
BORDA_CODIGO = Border(
    left=Side(style="medium"),
    right=Side(style="medium"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

COLUNAS_FIBRAS = (
    ("COD CHASSIS", 28.43, 40, FONTE_CODIGO_BARRAS),
    ("POP ID", 21.00, 28, FONTE_PADRAO),
    ("CHASSI", 8.00, 10, FONTE_PADRAO),
    ("MODELO", 12.43, 8, FONTE_PADRAO),
    ("MU", 5.00, 8, FONTE_PADRAO),
    ("COR", 6.57, 8, FONTE_PADRAO),
    ("DEF", 4.86, 8, FONTE_PADRAO),
    ("LATERAL", 8.14, 8, FONTE_PADRAO),
    ("GRELHA", 13.86, 6, FONTE_PADRAO),
)


COLUNAS_PADRAO = tuple(nome for nome, *_ in COLUNAS_FIBRAS)


def _chave_coluna(nome) -> str:
    return " ".join(str(nome).strip().upper().split())


def _colunas_escolhidas(colunas_desejadas=None):
    """
    Filtra o layout pelas colunas pedidas, mantendo a ordem original.

    O layout desta funcao e fixo (largura, fonte, codigo de barras por coluna),
    entao configurar significa escolher quais dessas colunas entram — nao
    inventar colunas novas.
    """
    if not colunas_desejadas:
        return COLUNAS_FIBRAS
    pedidas = {_chave_coluna(c) for c in colunas_desejadas}
    escolhidas = tuple(c for c in COLUNAS_FIBRAS if _chave_coluna(c[0]) in pedidas)
    return escolhidas or COLUNAS_FIBRAS


def _fill(cor: str) -> PatternFill:
    return PatternFill(start_color=cor, end_color=cor, fill_type="solid")


def _normalizar_colunas(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df.columns = [
        " ".join(str(coluna or "").strip().upper().split())
        for coluna in df.columns
    ]
    return df


def _coluna(df: pd.DataFrame, *opcoes: str, obrigatoria: bool = True) -> pd.Series:
    for nome in opcoes:
        if nome in df.columns:
            return df[nome]

    if obrigatoria:
        raise ValueError(
            "O MIX geral não contém a coluna obrigatória: "
            + " ou ".join(opcoes)
            + "."
        )
    return pd.Series([pd.NA] * len(df), index=df.index, dtype="object")


def _valor_preenchido(valor: object) -> bool:
    return (
        pd.notna(valor)
        and str(valor).strip() != ""
        and str(valor).strip().lower() != "nan"
    )


def _gerar_grelha(df: pd.DataFrame) -> pd.Series:
    """Deriva somente os estados que existem objetivamente no MIX universal."""
    if "GRELHA" in df.columns:
        return df["GRELHA"]

    resultado = pd.Series([pd.NA] * len(df), index=df.index, dtype="object")
    level = _coluna(df, "LEVEL", obrigatoria=False)
    climatizador = _coluna(df, "AUX CAB COOLER", obrigatoria=False)

    resultado.loc[climatizador.apply(_valor_preenchido)] = "CLIMATIZADOR"
    resultado.loc[level.fillna("").astype(str).str.strip().str.upper().eq("KD")] = "KD"
    return resultado


def preparar_mix_fibras(caminho_mix_geral: str) -> pd.DataFrame:
    """Seleciona do MIX geral apenas as colunas usadas pelo MIX DE FIBRAS."""
    df_geral = _normalizar_colunas(pd.read_excel(caminho_mix_geral, sheet_name=0))

    mascara_hide = df_geral.astype(str).apply(
        lambda linha: linha.str.contains("Hide", case=False, na=False).any(),
        axis=1,
    )
    df_geral = df_geral[~mascara_hide].copy()

    df = pd.DataFrame(
        {
            "POP ID": _coluna(df_geral, "POPID", "POP ID"),
            "CHASSI": _coluna(df_geral, "CHASSI"),
            "MODELO": _coluna(df_geral, "MODELO"),
            "MU": _coluna(df_geral, "MU"),
            "COR": _coluna(df_geral, "COR"),
            "DEF": _coluna(df_geral, "DEF"),
            "LATERAL": _coluna(df_geral, "LATERAL"),
            "GRELHA": _gerar_grelha(df_geral),
        }
    )

    return (
        df.dropna(subset=["POP ID", "CHASSI"], how="all")
        .reset_index(drop=True)
    )


def _codigo_barras_chassi(valor: object) -> str:
    if pd.isna(valor):
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = str(valor).strip()
    if texto.startswith("*") and texto.endswith("*"):
        return texto
    return f"*{texto}*"


def anexar_mix_fibras(
    caminho_mix_geral: str,
    caminho_arquivo_formatado: str,
    data_referencia: date | None = None,
    colunas_desejadas=None,
) -> None:
    """Cria no arquivo final a aba MIX DE FIBRAS, sem usar a MONA."""
    df = preparar_mix_fibras(caminho_mix_geral)
    data_referencia = data_referencia or date.today()
    nome_aba = f"MIX DE FIBRAS {data_referencia.strftime('%d-%m-%Y')}"

    wb = load_workbook(caminho_arquivo_formatado)
    try:
        if nome_aba in wb.sheetnames:
            del wb[nome_aba]
        ws = wb.create_sheet(nome_aba)

        colunas = _colunas_escolhidas(colunas_desejadas)
        indice_grelha = next(
            (i for i, c in enumerate(colunas, start=1) if c[0] == "GRELHA"), None)
        indice_codigo = next(
            (i for i, c in enumerate(colunas, start=1) if c[0] == "COD CHASSIS"), None)
        indice_def = next(
            (i for i, c in enumerate(colunas, start=1) if c[0] == "DEF"), None)
        indice_popid = next(
            (i for i, c in enumerate(colunas, start=1) if c[0] == "POP ID"), None)
        indice_chassi = next(
            (i for i, c in enumerate(colunas, start=1) if c[0] == "CHASSI"), None)

        for coluna, (titulo, largura, tamanho, fonte) in enumerate(
            colunas,
            start=1,
        ):
            celula = ws.cell(row=1, column=coluna, value=titulo)
            celula.font = Font(
                name=fonte,
                size=25 if titulo == "COD CHASSIS" else tamanho,
                bold=(titulo != "COD CHASSIS"),
                color="FFFFFFFF" if titulo == "GRELHA" else "FF000000",
            )
            celula.fill = _fill(AZUL if titulo == "GRELHA" else CINZA_CABECALHO)
            celula.alignment = CENTRO
            celula.border = BORDA_CODIGO if coluna == indice_codigo else BORDA
            ws.column_dimensions[get_column_letter(coluna)].width = largura

        ws.row_dimensions[1].height = 31.5

        for indice, linha in df.iterrows():
            numero_linha = indice + 2
            ws.row_dimensions[numero_linha].height = 39.95
            por_titulo = {
                "COD CHASSIS": _codigo_barras_chassi(linha["CHASSI"]),
                "POP ID": linha["POP ID"],
                "CHASSI": linha["CHASSI"],
                "MODELO": linha["MODELO"],
                "MU": linha["MU"],
                "COR": linha["COR"],
                "DEF": linha["DEF"],
                "LATERAL": linha["LATERAL"],
                "GRELHA": linha["GRELHA"],
            }
            valores = [por_titulo[titulo] for titulo, *_ in colunas]

            for coluna, valor in enumerate(valores, start=1):
                if pd.isna(valor):
                    valor = None
                celula = ws.cell(numero_linha, coluna, valor)
                celula.alignment = CENTRO
                celula.border = BORDA_CODIGO if coluna == indice_codigo else BORDA

                if coluna == indice_codigo:
                    celula.font = Font(
                        name=FONTE_CODIGO_BARRAS,
                        size=40,
                        bold=False,
                    )
                elif coluna == indice_popid:
                    celula.font = Font(name=FONTE_PADRAO, size=28, bold=True)
                    celula.number_format = "0"
                else:
                    tamanho = (6 if coluna == indice_grelha
                               else (10 if coluna == indice_chassi else 8))
                    celula.font = Font(
                        name=FONTE_PADRAO,
                        size=tamanho,
                        bold=(coluna == indice_grelha),
                    )

            valor_def = "" if pd.isna(linha["DEF"]) else str(linha["DEF"])
            if indice_def and valor_def.strip().upper() == "TETO":
                ws.cell(numero_linha, indice_def).fill = _fill(AMARELO)

            valor_grelha = (
                "" if pd.isna(linha["GRELHA"]) else str(linha["GRELHA"])
            ).strip().upper()
            estilos_grelha = {
                "KD": (VERDE_CLARO, "FF000000"),
                "CLIMATIZADOR": (AMARELO, "FF000000"),
                "PÓS TRANSPORTE": (VERDE, "FFFFFFFF"),
                "TVL": (AZUL, "FFFFFFFF"),
                "REALOCADO": ("FFDAEEF3", VERDE),
            }
            if indice_grelha and valor_grelha in estilos_grelha:
                fundo, fonte = estilos_grelha[valor_grelha]
                celula = ws.cell(numero_linha, indice_grelha)
                celula.fill = _fill(fundo)
                celula.font = Font(
                    name=FONTE_PADRAO,
                    size=6,
                    bold=True,
                    color=fonte,
                )

        ultima = get_column_letter(max(len(colunas), 1))
        ws.print_area = f"A1:{ultima}{max(ws.max_row, 1)}"
        ws.print_title_rows = "1:1"
        ws.freeze_panes = "A2"
        ws.page_setup.orientation = "portrait"
        ws.page_margins.left = ws.page_margins.right = 0.118
        ws.page_margins.top = ws.page_margins.bottom = 0.157
        wb.save(caminho_arquivo_formatado)
    finally:
        wb.close()


def formatar_excel_mix_fibras_logic(
    caminho_mix_geral: str,
    caminho_arquivo_novo: str,
    colunas_desejadas=None,
):
    """Entrada independente no mesmo padrão dos demais processadores."""
    try:
        from openpyxl import Workbook

        wb = Workbook()
        wb.save(caminho_arquivo_novo)
        wb.close()
        anexar_mix_fibras(caminho_mix_geral, caminho_arquivo_novo,
                          colunas_desejadas=colunas_desejadas)

        wb = load_workbook(caminho_arquivo_novo)
        if "Sheet" in wb.sheetnames and len(wb.sheetnames) > 1:
            del wb["Sheet"]
        wb.save(caminho_arquivo_novo)
        wb.close()
        return True, "MIX de fibras formatado com sucesso!"
    except Exception as erro:
        return False, f"Erro ao formatar o MIX de fibras: {erro}"
